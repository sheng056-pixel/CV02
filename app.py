import os
import io
import re
import hashlib
import pickle
import base64
import urllib.parse
from datetime import datetime, date
import requests
import pandas as pd
import plotly.express as px
import streamlit as st
import streamlit.components.v1 as components
from google import genai
from google.genai import types
import pypdf
import docx
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import qn, nsdecls
from pptx import Presentation
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity

# ================= 1. 系統基本設定與緊湊型 UI 樣式 =================
st.set_page_config(page_title="CV02 工程助手", page_icon="🏗️", layout="wide")

st.markdown("""
<style>
div[data-testid="stSidebar"] div.stButton > button {
    padding-top: 0.15rem;
    padding-bottom: 0.15rem;
    font-size: 0.85rem;
    text-align: left;
    justify-content: flex-start;
}
</style>
""", unsafe_allow_html=True)

TEAM_PASSWORD = st.secrets.get("TEAM_PASSWORD", "1234")
GEMINI_API_KEY = st.secrets.get("GEMINI_API_KEY", "").strip()
GITHUB_TOKEN = st.secrets.get("GITHUB_TOKEN", "").strip()
GITHUB_REPO = st.secrets.get("GITHUB_REPO", "").strip()
DB_FILE = "engineering_knowledge_db.pkl"

# ★ 已修復多模型自動備援清單（優先使用最穩定快速的官方模型）
AVAILABLE_MODELS = ["gemini-2.5-flash", "gemini-2.5-pro", "gemini-2.0-flash", "gemini-3.5-flash-lite", "gemini-3.5-flash"]

DEFAULT_FOLDER_TREE_V7 = {
    "01_招標文件與契約": {
        "第一次公告版本": ["契約條款", "投標與評選須知", "施工規範與圖說"],
        "第二次公告版本": ["契約條款", "投標與評選須知", "補充規定與附件"]
    },
    "02_施工規範與圖說": {
        "土建與水工": ["一般規範", "設計圖說"],
        "機電與弱電": ["設備規範", "界面圖說"],
        "職安衛與環保": ["職安衛規定", "環境保護計畫"]
    },
    "03_設計變更與核定公文": {
        "契約變更設計": ["變更設計圖說", "新增單價議價"],
        "重要往來公文": ["業主來文", "監造與承商發文"]
    },
    "04_會議紀錄與簡報": {
        "專案進度簡報": ["週會與月會簡報", "專案議題討論會"],
        "工務會議紀錄": ["歷次會議紀錄", "現勘結論"]
    },
    "05_材料送審與查驗紀錄": {
        "施工日報與查驗": ["施工日報表", "現場自主檢查表"],
        "材料送審資料": ["材料送審核定", "試驗報告"]
    }
}

DEFAULT_USERS = {
    "聖引": {"password": TEAM_PASSWORD, "role": "管理員", "display_name": "聖引"},
    "668436": {"password": TEAM_PASSWORD, "role": "一般組員", "display_name": "淳凱"},
    "077343": {"password": TEAM_PASSWORD, "role": "一般組員", "display_name": "有朋"},
    "479701": {"password": TEAM_PASSWORD, "role": "一般組員", "display_name": "旭暉"},
    "331165": {"password": TEAM_PASSWORD, "role": "一般組員", "display_name": "宣頷"},
    "組員A": {"password": TEAM_PASSWORD, "role": "一般組員", "display_name": "組員A"}
}

DEFAULT_SHORTCUTS = [
    "📌 最新進度與預算",
    "⏳ 關鍵要徑與浮時",
    "📅 鋼筋工項施工日期統計",
    "📋 歷次會議決議追蹤",
    "⚠️ 職安衛查驗標準"
]

DEFAULT_CLOUD_LINKS = [
    {"title": "📁 每日施工照片雲端夾", "url": "https://drive.google.com/drive/folders/1MFrIjPZB3cBWBb2Nt_q2eaTRDw7zzno9?usp=drive_link"},
    {"title": "📁 每日日報更新雲端夾", "url": "https://drive.google.com/drive/folders/1vmmaXftLfCTT3fd5_x_sIl6jc9KrEATX?usp=drive_link"}
]

DEFAULT_FIGMA_BOARDS = []

# ★ 預設甘特圖排程資料（隧道明挖工法標準流程）
DEFAULT_GANTT_DATA = [
    {"項次": "1", "WBS": "1.0", "工項名稱": "施工前置與測量放樣", "預計開始日": "2026-11-01", "預計完成日": "2026-11-10", "工期(天)": 10, "進度(%)": 100, "狀態": "已完成", "是否要徑": False, "備註": "現地勘查與管線遷移"},
    {"項次": "2", "WBS": "2.0", "工項名稱": "打全套管基樁 / 深基礎", "預計開始日": "2026-11-11", "預計完成日": "2026-12-10", "工期(天)": 30, "進度(%)": 60, "狀態": "進行中", "是否要徑": True, "備註": "兩側深基礎壁體"},
    {"項次": "3", "WBS": "3.0", "工項名稱": "打鋼板樁（或連續壁）", "預計開始日": "2026-12-11", "預計完成日": "2026-12-31", "工期(天)": 21, "進度(%)": 0, "狀態": "未開始", "是否要徑": True, "備註": "防止土石崩塌與滲水"},
    {"項次": "4", "WBS": "4.1", "工項名稱": "分層開挖（第一階）", "預計開始日": "2027-01-04", "預計完成日": "2027-01-18", "工期(天)": 15, "進度(%)": 0, "狀態": "未開始", "是否要徑": True, "備註": "配合開挖深度向下"},
    {"項次": "5", "WBS": "4.2", "工項名稱": "搭設擋土支撐（H型鋼）", "預計開始日": "2027-01-19", "預計完成日": "2027-01-28", "工期(天)": 10, "進度(%)": 0, "狀態": "未開始", "是否要徑": False, "備註": "確保壁面穩定"},
    {"項次": "6", "WBS": "5.1", "工項名稱": "繼續向下開挖至底面", "預計開始日": "2027-01-29", "預計完成日": "2027-02-15", "工期(天)": 18, "進度(%)": 0, "狀態": "未開始", "是否要徑": True, "備註": "挖至設計基礎底面"},
    {"項次": "7", "WBS": "5.2", "工項名稱": "打底層混凝土（PC）", "預計開始日": "2027-02-16", "預計完成日": "2027-02-20", "工期(天)": 5, "進度(%)": 0, "狀態": "未開始", "是否要徑": False, "備註": "提供平整工作面"},
    {"項次": "8", "WBS": "6.1", "工項名稱": "綁紮基礎鋼筋與模板組立", "預計開始日": "2027-02-21", "預計完成日": "2027-03-05", "工期(天)": 13, "進度(%)": 0, "狀態": "未開始", "是否要徑": True, "備註": "配置底版鋼筋與剪力筋"},
    {"項次": "9", "WBS": "6.2", "工項名稱": "基礎版混凝土澆置與養護", "預計開始日": "2027-03-06", "預計完成日": "2027-03-12", "工期(天)": 7, "進度(%)": 0, "狀態": "未開始", "是否要徑": True, "備註": "大宗混凝土澆置"},
    {"項次": "10", "WBS": "7.0", "工項名稱": "結構體主體施工（牆與頂版）", "預計開始日": "2027-03-13", "預計完成日": "2027-05-15", "工期(天)": 64, "進度(%)": 0, "狀態": "未開始", "是否要徑": True, "備註": "側牆、中柱與頂版"},
    {"項次": "11", "WBS": "8.0", "工項名稱": "防水工程施作", "預計開始日": "2027-05-16", "預計完成日": "2027-05-31", "工期(天)": 16, "進度(%)": 0, "狀態": "未開始", "是否要徑": False, "備註": "外壁及頂版防水膜"},
    {"項次": "12", "WBS": "9.0", "工項名稱": "分層回填與地表復原", "預計開始日": "2027-06-01", "預計完成日": "2027-06-20", "工期(天)": 20, "進度(%)": 0, "狀態": "未開始", "是否要徑": False, "備註": "分層夯實並恢復設施"}
]

ENGINEERING_SYNONYMS = {
    "浮實": ["浮時", "總浮時", "自由浮時", "浮動", "Total Float", "TF"],
    "浮時": ["浮實", "總浮時", "自由浮時", "浮動", "Total Float", "TF"],
    "要徑": ["關鍵路徑", "關鍵要徑", "Critical Path", "CP", "進度", "落後", "趲趕"],
    "關鍵路徑": ["要徑", "關鍵要徑", "浮時"],
    "決議": ["結論", "裁示", "會議紀錄", "討論會", "拍板", "定案"],
    "預算": ["經費", "金額", "契約價金", "總價", "發包", "變更設計"],
    "施工日期": ["日報", "週報", "出工", "施作", "綁紮", "澆置", "開挖", "日曆天"],
    "里程碑": ["分項", "節點", "竣工", "期限", "工期"]
}

# ================= 自動 GitHub 同步備份機制 =================
def pull_db_from_github():
    if not GITHUB_TOKEN or not GITHUB_REPO:
        return None
    url = f"https://api.github.com/repos/{GITHUB_REPO}/contents/{DB_FILE}?ref=main"
    headers = {"Authorization": f"Bearer {GITHUB_TOKEN}", "Accept": "application/vnd.github.v3+json"}
    try:
        r = requests.get(url, headers=headers, timeout=10)
        if r.status_code == 200:
            content_b64 = r.json().get("content", "")
            if content_b64:
                return base64.b64decode(content_b64)
    except Exception:
        pass
    return None

def push_db_to_github(db_bytes):
    if not GITHUB_TOKEN or not GITHUB_REPO:
        return
    url = f"https://api.github.com/repos/{GITHUB_REPO}/contents/{DB_FILE}"
    headers = {"Authorization": f"Bearer {GITHUB_TOKEN}", "Accept": "application/vnd.github.v3+json"}
    try:
        r_get = requests.get(url, headers=headers, timeout=10)
        sha = r_get.json().get("sha", "") if r_get.status_code == 200 else ""
        data = {
            "message": "🔄 Auto-sync CV02 Database from Streamlit",
            "content": base64.b64encode(db_bytes).decode("utf-8"),
            "branch": "main"
        }
        if sha:
            data["sha"] = sha
        requests.put(url, headers=headers, json=data, timeout=15)
    except Exception:
        pass

def make_token(username, password):
    raw = f"{username}:{password}:cv02_salt"
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()[:16]

def ensure_three_level_tree(raw_tree):
    new_tree = {}
    for c1, val in raw_tree.items():
        if isinstance(val, dict):
            new_tree[c1] = {}
            for c2, c3_list in val.items():
                new_tree[c1][c2] = list(c3_list) if c3_list else ["預設細項"]
        elif isinstance(val, list):
            new_tree[c1] = {}
            for c2 in val:
                new_tree[c1][c2] = ["一般文件", "補充附件"]
    if not new_tree:
        return DEFAULT_FOLDER_TREE_V7.copy()
    return new_tree

def load_db():
    if not os.path.exists(DB_FILE) and GITHUB_TOKEN and GITHUB_REPO:
        remote_bytes = pull_db_from_github()
        if remote_bytes:
            with open(DB_FILE, "wb") as f:
                f.write(remote_bytes)

    if os.path.exists(DB_FILE):
        try:
            with open(DB_FILE, "rb") as f:
                data = pickle.load(f)
            if isinstance(data, dict):
                raw_tree = data.get("folder_tree", DEFAULT_FOLDER_TREE_V7.copy())
                tree_3d = ensure_three_level_tree(raw_tree)
                chunks = data.get("chunks", [])
                for c in chunks:
                    c1 = c.get("cat1", list(tree_3d.keys())[0])
                    if c1 not in tree_3d:
                        tree_3d[c1] = {"預設次夾": ["一般文件"]}
                    c2 = c.get("cat2", list(tree_3d[c1].keys())[0])
                    if c2 not in tree_3d[c1]:
                        tree_3d[c1][c2] = ["一般文件"]
                    if not tree_3d[c1][c2]:
                        tree_3d[c1][c2] = ["一般文件"]
                    c3 = c.get("cat3", "")
                    c["cat1"], c["cat2"], c["cat3"] = c1, c2, c3
                    if "uploader" not in c:
                        c["uploader"] = "聖引"
                return {
                    "folder_tree": tree_3d,
                    "chunks": chunks,
                    "users": data.get("users", DEFAULT_USERS.copy()),
                    "shortcuts": data.get("shortcuts", DEFAULT_SHORTCUTS.copy()),
                    "cloud_links": data.get("cloud_links", [dict(x) for x in DEFAULT_CLOUD_LINKS]),
                    "figma_boards": data.get("figma_boards", [dict(x) for x in DEFAULT_FIGMA_BOARDS]),
                    "gantt_data": data.get("gantt_data", [dict(x) for x in DEFAULT_GANTT_DATA]),
                    "gantt_title": data.get("gantt_title", "台電統包工程 — 隧道明挖工法施工進度甘特圖與排程管控表"),
                    "max_mb": data.get("max_mb", 80)
                }
        except Exception:
            pass
    return {
        "folder_tree": DEFAULT_FOLDER_TREE_V7.copy(),
        "chunks": [],
        "users": DEFAULT_USERS.copy(),
        "shortcuts": DEFAULT_SHORTCUTS.copy(),
        "cloud_links": [dict(x) for x in DEFAULT_CLOUD_LINKS],
        "figma_boards": [dict(x) for x in DEFAULT_FIGMA_BOARDS],
        "gantt_data": [dict(x) for x in DEFAULT_GANTT_DATA],
        "gantt_title": "台電統包工程 — 隧道明挖工法施工進度甘特圖與排程管控表",
        "max_mb": 80
    }

def save_db(db_dict):
    with open(DB_FILE, "wb") as f:
        pickle.dump(db_dict, f)
    if GITHUB_TOKEN and GITHUB_REPO:
        with open(DB_FILE, "rb") as f:
            push_db_to_github(f.read())

def to_figma_embed_url(raw_url):
    u = raw_url.strip()
    if "iframe" in u and "src=" in u:
        m = re.search(r'src=["\']([^"\']+)["\']', u)
        if m:
            return m.group(1)
    if "embed.figma.com" in u or "figma.com/embed" in u:
        return u
    if "figma.com/board/" in u or "figma.com/design/" in u or "figma.com/file/" in u or "figma.com/proto/" in u:
        embed_u = u.replace("https://www.figma.com/", "https://embed.figma.com/").replace("https://figma.com/", "https://embed.figma.com/")
        sep = "&" if "?" in embed_u else "?"
        if "embed-host=" not in embed_u:
            embed_u += f"{sep}embed-host=share"
        return embed_u
    encoded = urllib.parse.quote(u, safe="")
    return f"https://www.figma.com/embed?embed_host=share&url={encoded}"

# ================= 甘特圖 Excel 與開會列印報表匯出引擎 =================
def generate_gantt_excel_bytes(gantt_list, title_text):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "施工進度排程表"
    ws.views.sheetView[0].showGridLines = True

    ws.merge_cells("A1:J1")
    title_cell = ws["A1"]
    title_cell.value = title_text
    title_cell.font = Font(name="標楷體", size=16, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 35

    headers = ["項次", "WBS", "工項名稱", "預計開始日", "預計完成日", "工期(天)", "進度(%)", "狀態", "是否要徑", "備註與工法說明"]
    ws.append([])
    ws.append(headers)
    ws.row_dimensions[3].height = 25

    header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    header_font = Font(name="標楷體", size=11, bold=True, color="000000")
    thin_border = Border(
        left=Side(style='thin', color='BFBFBF'),
        right=Side(style='thin', color='BFBFBF'),
        top=Side(style='thin', color='BFBFBF'),
        bottom=Side(style='thin', color='BFBFBF')
    )

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=3, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    cp_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
    for row_idx, r in enumerate(gantt_list, start=4):
        is_cp = bool(r.get("是否要徑", False))
        ws.append([
            str(r.get("項次", "")),
            str(r.get("WBS", "")),
            str(r.get("工項名稱", "")),
            str(r.get("預計開始日", "")),
            str(r.get("預計完成日", "")),
            int(r.get("工期(天)", 1)),
            int(r.get("進度(%)", 0)),
            str(r.get("狀態", "未開始")),
            "🔥 是 (要徑)" if is_cp else "否",
            str(r.get("備註", ""))
        ])
        ws.row_dimensions[row_idx].height = 22
        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_num)
            cell.border = thin_border
            cell.font = Font(name="標楷體", size=10, bold=is_cp, color="C00000" if (is_cp and col_num in [3, 9]) else "000000")
            if is_cp:
                cell.fill = cp_fill
            if col_num in [1, 2, 4, 5, 8, 9]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_num in [6, 7]:
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
    ws.column_dimensions["C"].width = 32
    ws.column_dimensions["J"].width = 30

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

def generate_printable_gantt_html(gantt_list, title_text, fig_html):
    rows_html = ""
    for r in gantt_list:
        is_cp = bool(r.get("是否要徑", False))
        bg = "background-color: #fff2f0; font-weight: bold;" if is_cp else ""
        cp_badge = "<span style='color:#d93025;'>🔥 要徑</span>" if is_cp else "一般"
        rows_html += f"""
        <tr style="{bg}">
            <td style="text-align:center;">{r.get('項次','')}</td>
            <td style="text-align:center;">{r.get('WBS','')}</td>
            <td>{r.get('工項名稱','')}</td>
            <td style="text-align:center;">{r.get('預計開始日','')}</td>
            <td style="text-align:center;">{r.get('預計完成日','')}</td>
            <td style="text-align:right;">{r.get('工期(天)',0)} 天</td>
            <td style="text-align:right;">{r.get('進度(%)',0)}%</td>
            <td style="text-align:center;">{r.get('狀態','')}</td>
            <td style="text-align:center;">{cp_badge}</td>
            <td>{r.get('備註','')}</td>
        </tr>
        """
    html_str = f"""
    <!DOCTYPE html>
    <html>
    <head>
        <meta charset="utf-8">
        <title>{title_text}</title>
        <style>
            @page {{ size: A4 landscape; margin: 12mm; }}
            body {{ font-family: "Microsoft JhengHei", "標楷體", sans-serif; color: #222; margin: 20px; }}
            h2 {{ text-align: center; color: #1B365D; margin-bottom: 5px; }}
            .sub {{ text-align: center; font-size: 13px; color: #555; margin-bottom: 15px; }}
            .print-btn {{
                display: block; margin: 0 auto 15px auto; padding: 10px 24px;
                background: #d93025; color: white; border: none; border-radius: 6px;
                font-size: 15px; font-weight: bold; cursor: pointer;
            }}
            table {{ width: 100%; border-collapse: collapse; font-size: 12px; margin-top: 15px; }}
            th {{ background-color: #1B365D; color: white; padding: 7px; border: 1px solid #999; }}
            td {{ padding: 6px; border: 1px solid #ccc; }}
            @media print {{ .print-btn {{ display: none; }} }}
        </style>
    </head>
    <body>
        <button class="print-btn" onclick="window.print()">🖨️ 點此直接列印或另存為橫向 PDF (Ctrl+P)</button>
        <h2>{title_text}</h2>
        <div class="sub">製表日期：{date.today().strftime('%Y-%m-%d')} ｜ 🔥 標記代表關鍵要徑 (Critical Path)</div>
        <div>{fig_html}</div>
        <table>
            <thead>
                <tr>
                    <th>項次</th><th>WBS</th><th>工項名稱</th><th>預計開始日</th><th>預計完成日</th>
                    <th>工期</th><th>進度</th><th>狀態</th><th>屬性</th><th>備註說明</th>
                </tr>
            </thead>
            <tbody>{rows_html}</tbody>
        </table>
    </body>
    </html>
    """
    return html_str.encode("utf-8")

# ================= 台電標準「簽辦用紙（內簽）」紅框 Word 生成引擎 =================
def set_red_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        '<w:top w:val="single" w:sz="6" w:space="0" w:color="FF0000"/>'
        '<w:left w:val="single" w:sz="6" w:space="0" w:color="FF0000"/>'
        '<w:bottom w:val="single" w:sz="6" w:space="0" w:color="FF0000"/>'
        '<w:right w:val="single" w:sz="6" w:space="0" w:color="FF0000"/>'
        '<w:insideH w:val="single" w:sz="4" w:space="0" w:color="FF0000"/>'
        '<w:insideV w:val="single" w:sz="6" w:space="0" w:color="FF0000"/>'
        '</w:tblBorders>'
    )
    tblPr.append(borders)

def set_cell_TopBottom_borders(cell, top_sz="4", bottom_sz="4", left_sz="6", right_sz="6"):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="{top_sz}" w:space="0" w:color="FF0000"/>'
        f'<w:left w:val="single" w:sz="{left_sz}" w:space="0" w:color="FF0000"/>'
        f'<w:bottom w:val="single" w:sz="{bottom_sz}" w:space="0" w:color="FF0000"/>'
        f'<w:right w:val="single" w:sz="{right_sz}" w:space="0" w:color="FF0000"/>'
        '</w:tcBorders>'
    )
    tcPr.append(borders)

def format_run_font(run, font_name="標楷體", size_pt=12, is_red=False, bold=False):
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
    run.font.size = Pt(size_pt)
    run.bold = bold
    if is_red:
        run.font.color.rgb = RGBColor(255, 0, 0)
    else:
        run.font.color.rgb = RGBColor(0, 0, 0)

def generate_taipower_internal_memo_docx(unit_name, doc_no, file_no, retention, memo_text):
    doc = docx.Document()
    for section in doc.sections:
        section.top_margin = Cm(1.5)
        section.bottom_margin = Cm(1.5)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)

    header_tbl = doc.add_table(rows=3, cols=3)
    header_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_tbl.autofit = False

    col_widths_top = [Cm(10.5), Cm(2.2), Cm(4.3)]
    for row in header_tbl.rows:
        for idx, w in enumerate(col_widths_top):
            row.cells[idx].width = w

    title_cell = header_tbl.cell(0, 0).merge(header_tbl.cell(2, 0))
    title_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p_title = title_cell.paragraphs[0]
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_t1 = p_title.add_run("台灣電力公司 ")
    format_run_font(r_t1, "標楷體", 20, is_red=True, bold=True)
    r_t2 = p_title.add_run("簽辦用紙")
    format_run_font(r_t2, "標楷體", 18, is_red=True, bold=True)

    meta_labels = [("文  號", doc_no), ("檔  號", file_no), ("保存年限", retention)]
    for r_idx, (lbl, val) in enumerate(meta_labels):
        c_lbl = header_tbl.cell(r_idx, 1)
        c_val = header_tbl.cell(r_idx, 2)
        set_cell_TopBottom_borders(c_lbl)
        set_cell_TopBottom_borders(c_val)
        p1 = c_lbl.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.DISTRIBUTE
        r1 = p1.add_run(lbl)
        format_run_font(r1, "標楷體", 10, is_red=True)

        p2 = c_val.paragraphs[0]
        r2 = p2.add_run(val)
        format_run_font(r2, "標楷體", 10, is_red=False)

    main_tbl = doc.add_table(rows=13, cols=2)
    main_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    main_tbl.autofit = False
    set_red_table_borders(main_tbl)

    for row in main_tbl.rows:
        row.cells[0].width = Cm(2.5)
        row.cells[1].width = Cm(14.5)

    h_left = main_tbl.cell(0, 0).paragraphs[0]
    h_left.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_hl = h_left.add_run("簽辦單位")
    format_run_font(r_hl, "標楷體", 12, is_red=True)

    h_right = main_tbl.cell(0, 1).paragraphs[0]
    h_right.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_hr = h_right.add_run("擬  辦  或  會  核  意  見")
    format_run_font(r_hr, "標楷體", 12, is_red=True)

    u_p = main_tbl.cell(1, 0).paragraphs[0]
    u_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_unit = u_p.add_run(unit_name)
    format_run_font(r_unit, "標楷體", 14, is_red=False)

    content_cell = main_tbl.cell(1, 1).merge(main_tbl.cell(12, 1))
    first_p = True
    for raw_line in memo_text.split("\n"):
        clean_line = raw_line.replace("###", "").replace("**", "").strip()
        if not clean_line:
            continue
        p = content_cell.paragraphs[0] if first_p else content_cell.add_paragraph()
        first_p = False
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.35
        r_line = p.add_run(clean_line)
        format_run_font(r_line, "標楷體", 13, is_red=False)

    if "陳核" not in memo_text and "陳 核" not in memo_text:
        p_sign = content_cell.add_paragraph()
        p_sign.paragraph_format.space_before = Pt(18)
        r_sign = p_sign.add_run("  陳\n核")
        format_run_font(r_sign, "標楷體", 14, is_red=False)

    p_foot = doc.add_paragraph()
    p_foot.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_f1 = p_foot.add_run("第 ")
    format_run_font(r_f1, "標楷體", 10, is_red=True)
    r_f2 = p_foot.add_run("1 / 1")
    format_run_font(r_f2, "標楷體", 10, is_red=False)
    r_f3 = p_foot.add_run(" 頁")
    format_run_font(r_f3, "標楷體", 10, is_red=True)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()

db = load_db()
users_db = db["users"]
folder_tree = db["folder_tree"]
chunks = db["chunks"]
max_mb_limit = db.get("max_mb", 80)

# ================= AI 智慧自動分類路徑預測引擎 =================
def auto_classify_folder(filename, folder_tree):
    fn = filename.lower()
    for c1, c2_dict in folder_tree.items():
        for c2, c3_list in c2_dict.items():
            for c3 in c3_list:
                if c3.lower() in fn:
                    return c1, c2, c3
            c2_clean = re.sub(r"^\d+_", "", c2).lower()
            if c2_clean and c2_clean in fn:
                return c1, c2, (c3_list[0] if c3_list else "")

    rules = [
        (["日報", "自主檢查", "查驗"], "05_材料送審與查驗紀錄", "施工日報與查驗", "施工日報表"),
        (["送審", "試驗", "試體", "出廠", "型錄"], "05_材料送審與查驗紀錄", "材料送審資料", "材料送審核定"),
        (["會議", "現勘", "紀錄", "討論會"], "04_會議紀錄與簡報", "工務會議紀錄", "歷次會議紀錄"),
        (["簡報", "週會", "月會", "進度", "ppt"], "04_會議紀錄與簡報", "專案進度簡報", "週會與月會簡報"),
        (["變更", "議價", "新增單價", "修正契約"], "03_設計變更與核定公文", "契約變更設計", "變更設計圖說"),
        (["函", "公文", "來文", "發文", "備忘錄", "簽"], "03_設計變更與核定公文", "重要往來公文", "業主來文"),
        (["職安", "安衛", "環保", "安全"], "02_施工規範與圖說", "職安衛與環保", "職安衛規定"),
        (["機電", "弱電", "電氣", "空調", "消防"], "02_施工規範與圖說", "機電與弱電", "設備規範"),
        (["規範", "圖說", "施工", "結構", "土建", "水工"], "02_施工規範與圖說", "土建與水工", "一般規範"),
        (["契約", "投標", "招標", "須知", "價目表", "預算", "單價"], "01_招標文件與契約", "第二次公告版本", "契約條款"),
    ]
    for kws, target_c1, target_c2, target_c3 in rules:
        if any(k in fn for k in kws):
            if target_c1 in folder_tree:
                c2_dict = folder_tree[target_c1]
                c2_use = target_c2 if target_c2 in c2_dict else list(c2_dict.keys())[0]
                c3_list = c2_dict[c2_use]
                c3_use = target_c3 if target_c3 in c3_list else (c3_list[0] if c3_list else "")
                return target_c1, c2_use, c3_use

    def_c1 = list(folder_tree.keys())[0]
    def_c2 = list(folder_tree[def_c1].keys())[0]
    def_c3 = folder_tree[def_c1][def_c2][0] if folder_tree[def_c1][def_c2] else ""
    return def_c1, def_c2, def_c3

# ================= 登入驗證 =================
if "current_user" not in st.session_state:
    st.session_state.current_user = None

q_user = st.query_params.get("user", "")
q_token = st.query_params.get("token", "")

if not st.session_state.current_user and q_user in users_db:
    expected_token = make_token(q_user, users_db[q_user]["password"])
    if q_token == expected_token:
        st.session_state.current_user = q_user

if not st.session_state.current_user:
    st.title("🏗️ CV02 工程助手 — 登入")
    st.caption("勾選「記住我」後將網頁加入手機主畫面或書籤，下次點開即可直接進入系統！")
    with st.container(border=True):
        login_options = list(users_db.keys())
        def format_user_label(u):
            if u in users_db and "display_name" in users_db[u]:
                return f"{u} ({users_db[u]['display_name']})"
            return u

        login_user = st.selectbox("👤 選擇帳號 / 輸入工號：", login_options, format_func=format_user_label)
        login_pwd = st.text_input("🔑 輸入密碼：", type="password")
        remember_me = st.checkbox("📌 在此裝置記住我的登入狀態", value=True)
        if st.button("🚀 登入", type="primary", use_container_width=True):
            valid_pwd = users_db[login_user]["password"]
            if login_pwd == valid_pwd or login_pwd == TEAM_PASSWORD:
                st.session_state.current_user = login_user
                if remember_me:
                    st.query_params["user"] = login_user
                    st.query_params["token"] = make_token(login_user, valid_pwd)
                st.rerun()
            else:
                st.error("密碼錯誤！")
    st.stop()

raw_current_user = st.session_state.current_user
current_user = users_db.get(raw_current_user, {}).get("display_name", raw_current_user)

# ================= 檔案解析與匯入核心函數 =================
def ocr_with_gemini(file_bytes, mime_type):
    if not GEMINI_API_KEY:
        return ""
    client = genai.Client(api_key=GEMINI_API_KEY)
    prompt = (
        "請完整、精準地辨識並轉錄這份工程公文、契約、施工圖說、配筋表或照片中的所有繁體中文文字、文號、日期、表格數據與印章文字。"
        "若遇到工程配筋表或表格，請務必將『構材編號（如 G1, B1, C1, S1, W1）』與對應的『位置、主筋、箍筋、號數、間距、保護層、搭接長度』完整配對並以『欄位1 | 欄位2』格式呈現，不要遺漏任何細節。"
    )
    for model_name in AVAILABLE_MODELS:
        try:
            resp = client.models.generate_content(
                model=model_name,
                contents=[
                    types.Part.from_bytes(data=file_bytes, mime_type=mime_type),
                    prompt
                ],
                config=types.GenerateContentConfig(temperature=0.1)
            )
            if resp and resp.text:
                return resp.text.strip()
        except Exception:
            continue
    return ""

def extract_text_from_file(uploaded_file, cat1, cat2, cat3, uploader, enable_ocr=True):
    text_chunks = []
    filename = uploaded_file.name
    lower_name = filename.lower()
    raw_bytes = uploaded_file.getvalue()

    if lower_name.endswith((".jpg", ".jpeg", ".png")):
        mime = "image/png" if lower_name.endswith(".png") else "image/jpeg"
        ocr_text = ocr_with_gemini(raw_bytes, mime) if enable_ocr else ""
        if ocr_text:
            text_chunks.append({
                "file_name": filename, "cat1": cat1, "cat2": cat2, "cat3": cat3,
                "location": "(影像OCR辨識)", "content": ocr_text, "uploader": uploader
            })
        return text_chunks, (1 if ocr_text else 0)

    elif lower_name.endswith(".pdf"):
        reader = pypdf.PdfReader(io.BytesIO(raw_bytes))
        scan_pages = []
        for i, page in enumerate(reader.pages):
            try:
                page_text = (page.extract_text() or "").strip()
            except Exception:
                page_text = ""

            if len(page_text) >= 25:
                text_chunks.append({
                    "file_name": filename, "cat1": cat1, "cat2": cat2, "cat3": cat3,
                    "location": f"(第 {i+1} 頁)", "content": page_text, "uploader": uploader
                })
            else:
                scan_pages.append(i)

        ocr_count = 0
        if scan_pages and enable_ocr:
            max_ocr_pages = min(len(scan_pages), 25)
            prog = st.progress(0, text=f"🔍 偵測到《{filename}》含掃描頁或複雜工程圖說，正在啟動 AI 視覺辨識 (共 {max_ocr_pages} 頁)...")
            for idx, page_num in enumerate(scan_pages[:max_ocr_pages]):
                try:
                    writer = pypdf.PdfWriter()
                    writer.add_page(reader.pages[page_num])
                    buf = io.BytesIO()
                    writer.write(buf)
                    ocr_res = ocr_with_gemini(buf.getvalue(), "application/pdf")
                    if ocr_res:
                        text_chunks.append({
                            "file_name": filename, "cat1": cat1, "cat2": cat2, "cat3": cat3,
                            "location": f"(第 {page_num+1} 頁-AI圖說/掃描辨識)", "content": ocr_res, "uploader": uploader
                        })
                        ocr_count += 1
                except Exception:
                    pass
                prog.progress((idx + 1) / max_ocr_pages, text=f"🔍 正在辨識《{filename}》第 {page_num+1} 頁 ({idx+1}/{max_ocr_pages})...")
            prog.empty()
        return text_chunks, ocr_count

    elif lower_name.endswith(".docx"):
        doc = docx.Document(io.BytesIO(raw_bytes))
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            header_str = ""
            for r_idx, row in enumerate(table.rows):
                row_cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                row_text = " | ".join([c for c in row_cells if c])
                if not row_text:
                    continue
                if r_idx == 0:
                    header_str = row_text
                    paragraphs.append(f"【表格表頭：{header_str}】")
                else:
                    paragraphs.append(f"[{header_str}] ➔ {row_text}" if header_str else row_text)
        full_text = "\n".join(paragraphs)
        for i in range(0, len(full_text), 900):
            if full_text[i:i+900].strip():
                text_chunks.append({
                    "file_name": filename, "cat1": cat1, "cat2": cat2, "cat3": cat3,
                    "location": f"(段落 {i//900 + 1})", "content": full_text[i:i+900], "uploader": uploader
                })

    elif lower_name.endswith(".pptx"):
        prs = Presentation(io.BytesIO(raw_bytes))
        for i, slide in enumerate(prs.slides):
            slide_texts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        if paragraph.text.strip():
                            slide_texts.append(paragraph.text.strip())
                if shape.has_table:
                    header_str = ""
                    for r_idx, row in enumerate(shape.table.rows):
                        row_cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                        row_text = " | ".join([c for c in row_cells if c])
                        if not row_text:
                            continue
                        if r_idx == 0:
                            header_str = row_text
                            slide_texts.append(f"【表頭：{header_str}】")
                        else:
                            slide_texts.append(f"[{header_str}] ➔ {row_text}" if header_str else row_text)
            if slide_texts:
                text_chunks.append({
                    "file_name": filename, "cat1": cat1, "cat2": cat2, "cat3": cat3,
                    "location": f"(投影片第 {i+1} 頁)", "content": "\n".join(slide_texts), "uploader": uploader
                })

    elif lower_name.endswith(".xlsx"):
        wb = openpyxl.load_workbook(io.BytesIO(raw_bytes), data_only=True)
        for sheet in wb.sheetnames:
            ws = wb[sheet]
            rows_text = []
            header_str = ""
            for row in ws.iter_rows(values_only=True):
                row_vals = [str(cell).strip().replace("\n", " ") for cell in row if cell is not None and str(cell).strip()]
                if not row_vals:
                    continue
                line = " | ".join(row_vals)
                if not header_str:
                    header_str = line
                    rows_text.append(f"【工作表：{sheet}｜表頭欄位：{header_str}】")
                else:
                    rows_text.append(f"[{sheet}] {line}")
            for i in range(0, len(rows_text), 25):
                batch = rows_text[i:i+25]
                if i > 0 and header_str:
                    batch.insert(0, f"【表頭欄位：{header_str}】")
                chunk_str = "\n".join(batch)
                if chunk_str.strip():
                    text_chunks.append({
                        "file_name": filename, "cat1": cat1, "cat2": cat2, "cat3": cat3,
                        "location": f"(工作表:{sheet} 第{i+1}~{min(i+25, len(rows_text))}列)",
                        "content": chunk_str, "uploader": uploader
                    })

    elif lower_name.endswith(".txt"):
        full_text = raw_bytes.decode("utf-8", errors="ignore")
        for i in range(0, len(full_text), 900):
            if full_text[i:i+900].strip():
                text_chunks.append({
                    "file_name": filename, "cat1": cat1, "cat2": cat2, "cat3": cat3,
                    "location": f"(段落 {i//900 + 1})", "content": full_text[i:i+900], "uploader": uploader
                })

    return text_chunks, 0

def process_and_save_uploads(uploaded_files, mode="manual", target_c1="", target_c2="", target_c3="", enable_ocr=True):
    oversized = [uf.name for uf in uploaded_files if (uf.size / (1024 * 1024)) > max_mb_limit]
    img_count = sum(1 for uf in uploaded_files if uf.name.lower().endswith((".jpg", ".jpeg", ".png")))
    if oversized:
        st.error(f"⚠️ 超過 {max_mb_limit}MB 上限：\n" + "、".join(oversized))
        return False
    if img_count > 10:
        st.warning(f"⚠️ 一次選擇了 {img_count} 張圖片，建議單次 10 張以內！")
        return False

    with st.spinner("⚡ 正在極速解析內容並自動同步備份至 GitHub..."):
        imported_details = []
        total_ocr = 0
        for uf in uploaded_files:
            if mode == "smart":
                c1, c2, c3 = auto_classify_folder(uf.name, folder_tree)
            else:
                c1, c2, c3 = target_c1, target_c2, target_c3

            db["chunks"] = [c for c in db["chunks"] if c["file_name"] != uf.name]
            new_chunks, ocr_cnt = extract_text_from_file(uf, c1, c2, c3, current_user, enable_ocr=enable_ocr)
            if new_chunks:
                db["chunks"].extend(new_chunks)
                path_label = f"{c1} > {c2}" + (f" > {c3}" if c3 else " (第二層)")
                imported_details.append(f"`{uf.name}` ➔ 📂 **{path_label}**")
                total_ocr += ocr_cnt

        save_db(db)
        st.session_state.uploader_key += 1
        st.session_state.upload_notice = {
            "count": len(imported_details),
            "files": imported_details,
            "ocr_pages": total_ocr
        }
    return True

# ================= 彈出視窗 =================
@st.dialog("🎉 檔案極速匯入成功！")
def show_upload_success_dialog(info):
    st.success(f"✅ 已成功將 **{info['count']}** 份檔案寫入資料庫並自動同步至 GitHub！")
    if info.get("ocr_pages", 0) > 0:
        st.info(f"👁️ 內含 **{info['ocr_pages']}** 頁掃描影像或複雜圖說，已透過 AI 視覺辨識轉為可搜尋文字！")
    st.markdown("**📄 本次歸檔清單與路徑：**")
    for item_str in info.get("files", []):
        st.markdown(f"- {item_str}")
    if st.button("👌 太棒了，我知道了", type="primary", use_container_width=True):
        st.session_state.upload_notice = None
        st.rerun()

@st.dialog("📥 直接傳入此資料夾（拖拉放開即自動匯入）")
def show_quick_folder_upload_dialog(target_info):
    tc1, tc2, tc3 = target_info
    path_str = f"{tc1} > {tc2}" + (f" > {tc3}" if tc3 else " (直接放第二層)")
    st.markdown(f"📂 **目標鎖定路徑：** `{path_str}`")
    q_ocr = st.checkbox("👁️ 掃描檔與複雜圖說自動 AI 辨識 (OCR)", value=True, key="q_dlg_ocr")
    q_files = st.file_uploader(
        "請拖曳或選擇檔案：",
        type=["pdf", "docx", "pptx", "xlsx", "txt", "jpg", "jpeg", "png"],
        accept_multiple_files=True,
        key=f"q_dlg_up_{st.session_state.uploader_key}"
    )
    if q_files:
        if process_and_save_uploads(q_files, mode="manual", target_c1=tc1, target_c2=tc2, target_c3=tc3, enable_ocr=q_ocr):
            st.session_state.quick_upload_target = None
            st.rerun()
    if st.button("✖️ 取消關閉", use_container_width=True):
        st.session_state.quick_upload_target = None
        st.rerun()

@st.dialog("✏️ 編輯第三層細夾名稱")
def show_edit_c3_dialog(edit_info):
    ec1, ec2, ec3 = edit_info
    st.caption(f"📍 所屬位置：`{ec1} > {ec2}`")
    new_c3_val = st.text_input("✏️ 修改細夾名稱：", value=ec3)
    col_e1, col_e2, col_e3 = st.columns(3)
    with col_e1:
        if st.button("💾 儲存名稱", type="primary", use_container_width=True):
            clean_v = new_c3_val.strip()
            if clean_v and clean_v != ec3:
                idx = folder_tree[ec1][ec2].index(ec3)
                folder_tree[ec1][ec2][idx] = clean_v
                for c in chunks:
                    if c.get("cat1") == ec1 and c.get("cat2") == ec2 and c.get("cat3") == ec3:
                        c["cat3"] = clean_v
                save_db(db)
            st.session_state.edit_c3_target = None
            st.rerun()
    with col_e2:
        if st.button("🗑️ 刪除此細夾", use_container_width=True):
            if len(folder_tree[ec1][ec2]) > 1:
                folder_tree[ec1][ec2].remove(ec3)
                for c in chunks:
                    if c.get("cat1") == ec1 and c.get("cat2") == ec2 and c.get("cat3") == ec3:
                        c["cat3"] = ""
                save_db(db)
                st.session_state.edit_c3_target = None
                st.rerun()
            else:
                st.warning("至少需保留一個細夾！")
    with col_e3:
        if st.button("✖️ 取消", use_container_width=True):
            st.session_state.edit_c3_target = None
            st.rerun()

if "jump_to_file" not in st.session_state:
    st.session_state.jump_to_file = None
if "uploader_key" not in st.session_state:
    st.session_state.uploader_key = 0
if "quick_upload_target" not in st.session_state:
    st.session_state.quick_upload_target = None
if "edit_c3_target" not in st.session_state:
    st.session_state.edit_c3_target = None
if "last_chat_files" not in st.session_state:
    st.session_state.last_chat_files = []

if st.session_state.get("upload_notice"):
    show_upload_success_dialog(st.session_state.upload_notice)
elif st.session_state.get("quick_upload_target"):
    show_quick_folder_upload_dialog(st.session_state.quick_upload_target)
elif st.session_state.get("edit_c3_target"):
    show_edit_c3_dialog(st.session_state.edit_c3_target)

# ================= 6. 全工項動態智慧聯想引擎 =================
def expand_query_with_synonyms(query):
    expanded = query
    for key, syns in ENGINEERING_SYNONYMS.items():
        if key in query:
            expanded += " " + " ".join(syns)
    
    if GEMINI_API_KEY and len(query.strip()) <= 20:
        try:
            client = genai.Client(api_key=GEMINI_API_KEY)
            res = client.models.generate_content(
                model=AVAILABLE_MODELS[0],
                contents=f"請針對工程合約與詳細價目表中的查詢詞『{query}』，列出 3 至 5 個常見的工程同義詞、細部規格或相關工項名稱。請直接用空白隔開輸出這些詞彙，不要輸出其他多餘文字。",
                config=types.GenerateContentConfig(temperature=0.1)
            )
            if res and res.text:
                expanded += " " + res.text.strip().replace("\n", " ")
        except Exception:
            pass
    return expanded

def search_internal_db(query, chunks, selected_cat1s, selected_cat2s, selected_cat3s):
    if not chunks:
        return "（目前內部目錄尚無檔案，請先從左側上傳檔案）"

    filtered_chunks = [
        c for c in chunks
        if c.get("cat1") in selected_cat1s
        and (not selected_cat2s or c.get("cat2") in selected_cat2s)
        and (not selected_cat3s or (not c.get("cat3") or c.get("cat3") in selected_cat3s))
    ]
    if not filtered_chunks:
        return "（目前勾選的資料夾範圍內尚無文件）"

    deep_keywords = ["日期", "天數", "幾天", "所有", "全部", "統計", "歷次", "決議", "何時", "什麼時候", "要徑", "浮時", "浮實", "日報", "里程碑", "分項", "價目表", "工項", "單價", "拔除", "管"]
    is_deep_search = any(k in query for k in deep_keywords)
    top_k = 120 if is_deep_search else 25

    expanded_query = expand_query_with_synonyms(query)
    corpus = [f"{item['file_name']} {item['cat2']} {item.get('cat3', '')} {item['content']}" for item in filtered_chunks]
    vectorizer = TfidfVectorizer(analyzer="char_wb", ngram_range=(2, 4))
    tfidf_matrix = vectorizer.fit_transform(corpus + [expanded_query])
    cosine_sim = cosine_similarity(tfidf_matrix[-1], tfidf_matrix[:-1]).flatten()
    top_indices = cosine_sim.argsort()[-top_k:][::-1]

    selected_items = []
    for idx in top_indices:
        if cosine_sim[idx] > 0.005:
            selected_items.append(filtered_chunks[idx])

    if not selected_items:
        selected_items = filtered_chunks[:min(15, len(filtered_chunks))]

    if is_deep_search:
        selected_items.sort(key=lambda x: (x["file_name"], x["location"]))

    results = [
        f"【目錄：{item['cat1']} > {item['cat2']}{' > ' + item['cat3'] if item.get('cat3') else ''}｜來源：{item['file_name']} {item['location']}】\n{item['content']}"
        for item in selected_items
    ]
    return "\n\n---\n\n".join(results)

def highlight_keywords(text, keywords):
    highlighted = text
    for kw in keywords:
        if kw:
            pattern = re.compile(re.escape(kw), re.IGNORECASE)
            highlighted = pattern.sub(
                lambda m: f"<mark style='background-color: #ffe066; padding: 2px 4px; border-radius: 4px; font-weight: bold;'>{m.group(0)}</mark>",
                highlighted
            )
    return highlighted

file_info_map = {}
for c in chunks:
    c1 = c.get("cat1", list(folder_tree.keys())[0])
    if c1 not in folder_tree:
        folder_tree[c1] = {"預設次夾": ["一般文件"]}
    c2 = c.get("cat2", list(folder_tree[c1].keys())[0])
    if c2 not in folder_tree[c1]:
        folder_tree[c1][c2] = ["一般文件"]
    c3 = c.get("cat3", "")
    file_info_map[c["file_name"]] = (c1, c2, c3, c.get("uploader", "系統"))

db_size_mb = os.path.getsize(DB_FILE) / (1024 * 1024) if os.path.exists(DB_FILE) else 0.0

# ================= 左側導覽 =================
with st.sidebar:
    col_u1, col_u2 = st.columns([3, 2])
    with col_u1:
        st.markdown(f"👤 **{current_user}**")
    with col_u2:
        if st.button("🚪 登出", use_container_width=True):
            st.session_state.current_user = None
            st.query_params.clear()
            st.rerun()

    st.caption(f"📊 檔案：**{len(file_info_map)}** 份 ｜ 容量：**{db_size_mb:.2f} MB**")

    with st.expander("⚡ 智慧極速上傳 (AI 自動歸檔)", expanded=True):
        st.caption("💡 直接將多份檔案拖進下方框內，**手放開即自動依檔名分類並存入 GitHub**！")
        smart_ocr = st.checkbox("👁️ 掃描檔與圖說自動 AI 辨識 (OCR)", value=True, key="smart_ocr_chk")
        smart_files = st.file_uploader(
            f"拖拉檔案至此 (單檔上限 {max_mb_limit}MB)",
            type=["pdf", "docx", "pptx", "xlsx", "txt", "jpg", "jpeg", "png"],
            accept_multiple_files=True,
            key=f"smart_up_{st.session_state.uploader_key}"
        )
        if smart_files:
            if process_and_save_uploads(smart_files, mode="smart", enable_ocr=smart_ocr):
                st.rerun()

    with st.expander("➕ 手動指定資料夾上傳", expanded=False):
        up_c1 = st.selectbox("1. 第一層 (主夾)：", list(folder_tree.keys()))
        up_c2 = st.selectbox("2. 第二層 (次夾)：", list(folder_tree[up_c1].keys()))
        raw_c3_opts = folder_tree[up_c1][up_c2]
        c3_opts = ["(無細夾 /直接放第二層)"] + (raw_c3_opts if isinstance(raw_c3_opts, list) else [])
        up_c3_selected = st.selectbox("3. 第三層 (細夾，可選無)：", c3_opts)
        up_c3 = "" if up_c3_selected == "(無細夾 /直接放第二層)" else up_c3_selected

        c_chk1, c_chk2 = st.columns(2)
        with c_chk1:
            enable_ocr_upload = st.checkbox("👁️ 自動 OCR", value=True)
        with c_chk2:
            auto_submit_manual = st.checkbox("⚡ 選完即傳", value=True)

        uploaded_files = st.file_uploader(
            f"4. 選擇檔案 (上限 {max_mb_limit}MB)",
            type=["pdf", "docx", "pptx", "xlsx", "txt", "jpg", "jpeg", "png"],
            accept_multiple_files=True,
            key=f"file_up_{st.session_state.uploader_key}"
        )
        if uploaded_files and auto_submit_manual:
            if process_and_save_uploads(uploaded_files, mode="manual", target_c1=up_c1, target_c2=up_c2, target_c3=up_c3, enable_ocr=enable_ocr_upload):
                st.rerun()
        elif not auto_submit_manual:
            if st.button("🚀 匯入資料庫", type="primary", use_container_width=True):
                if not uploaded_files:
                    st.warning("請先選擇檔案！")
                else:
                    if process_and_save_uploads(uploaded_files, mode="manual", target_c1=up_c1, target_c2=up_c2, target_c3=up_c3, enable_ocr=enable_ocr_upload):
                        st.rerun()

    with st.expander("📁 目錄 (展開可直接改名 / 點 📥 快傳)", expanded=True):
        c_top1, c_top2 = st.columns([3, 2])
        with c_top1:
            show_empty_dirs = st.checkbox("顯示空資料夾", value=True)
        with c_top2:
            with st.popover("➕ 新增主夾", use_container_width=True):
                quick_new_c1 = st.text_input("輸入新第1層主夾名稱：", key="q_add_c1")
                if st.button("確認新增", key="btn_q_add_c1", type="primary", use_container_width=True) and quick_new_c1.strip():
                    if quick_new_c1.strip() not in folder_tree:
                        folder_tree[quick_new_c1.strip()] = {"預設次夾": ["一般文件"]}
                        save_db(db)
                        st.rerun()

        for c1, c2_dict in list(folder_tree.items()):
            c1_files = [f for f, (fc1, _, _, _) in file_info_map.items() if fc1 == c1]
            if not show_empty_dirs and not c1_files:
                continue
            with st.expander(f"🗂️ {c1} ({len(c1_files)})", expanded=False):
                with st.popover(f"✏️ 管理「{c1}」", use_container_width=True):
                    ren_c1_val = st.text_input("✏️ 修改第1層主夾名稱：", value=c1, key=f"in_ren_c1_{c1}")
                    if st.button("💾 儲存主夾名稱", key=f"btn_ren_c1_{c1}", type="primary", use_container_width=True):
                        clean_c1 = ren_c1_val.strip()
                        if clean_c1 and clean_c1 != c1:
                            db["folder_tree"] = {(clean_c1 if k == c1 else k): v for k, v in folder_tree.items()}
                            for c in chunks:
                                if c.get("cat1") == c1:
                                    c["cat1"] = clean_c1
                            save_db(db)
                            st.rerun()
                    st.divider()
                    add_c2_val = st.text_input("➕ 新增「第2層次夾」：", key=f"in_add_c2_{c1}")
                    if st.button("➕ 新增次夾", key=f"btn_add_c2_{c1}", use_container_width=True) and add_c2_val.strip():
                        if add_c2_val.strip() not in folder_tree[c1]:
                            folder_tree[c1][add_c2_val.strip()] = ["一般文件"]
                            save_db(db)
                            st.rerun()
                    st.divider()
                    if st.button("🗑️ 刪除此主夾", key=f"btn_del_c1_{c1}", use_container_width=True):
                        if len(folder_tree) > 1:
                            del folder_tree[c1]
                            fb1 = list(folder_tree.keys())[0]
                            fb2 = list(folder_tree[fb1].keys())[0]
                            for c in chunks:
                                if c.get("cat1") == c1:
                                    c["cat1"], c["cat2"], c["cat3"] = fb1, fb2, ""
                            save_db(db)
                            st.rerun()

                for c2, c3_list in list(c2_dict.items()):
                    c2_files = [f for f, (fc1, fc2, _, _) in file_info_map.items() if fc1 == c1 and fc2 == c2]
                    if not show_empty_dirs and not c2_files:
                        continue
                    with st.expander(f"📂 {c2} ({len(c2_files)})", expanded=False):
                        with st.popover(f"✏️ 管理「{c2}」", use_container_width=True):
                            ren_c2_val = st.text_input("✏️ 修改第2層次夾名稱：", value=c2, key=f"in_ren_c2_{c1}_{c2}")
                            if st.button("💾 儲存次夾名稱", key=f"btn_ren_c2_{c1}_{c2}", type="primary", use_container_width=True):
                                clean_c2 = ren_c2_val.strip()
                                if clean_c2 and clean_c2 != c2:
                                    folder_tree[c1] = {(clean_c2 if k == c2 else k): v for k, v in folder_tree[c1].items()}
                                    for c in chunks:
                                        if c.get("cat1") == c1 and c.get("cat2") == c2:
                                            c["cat2"] = clean_c2
                                    save_db(db)
                                    st.rerun()
                            st.divider()
                            add_c3_val = st.text_input("➕ 新增「第3層細夾」：", key=f"in_add_c3_{c1}_{c2}")
                            if st.button("➕ 新增細夾", key=f"btn_add_c3_{c1}_{c2}", use_container_width=True) and add_c3_val.strip():
                                if add_c3_val.strip() not in folder_tree[c1][c2]:
                                    folder_tree[c1][c2].append(add_c3_val.strip())
                                    save_db(db)
                                    st.rerun()
                            st.divider()
                            if st.button("🗑️ 刪除此次夾", key=f"btn_del_c2_{c1}_{c2}", use_container_width=True):
                                if len(folder_tree[c1]) > 1:
                                    del folder_tree[c1][c2]
                                    fb2 = list(folder_tree[c1].keys())[0]
                                    for c in chunks:
                                        if c.get("cat1") == c1 and c.get("cat2") == c2:
                                            c["cat2"], c["cat3"] = fb2, ""
                                    save_db(db)
                                    st.rerun()

                        col_d2_title, col_d2_btn = st.columns([4, 1])
                        with col_d2_title:
                            no_c3_files = sorted([f for f, (fc1, fc2, fc3, _) in file_info_map.items() if fc1 == c1 and fc2 == c2 and not fc3])
                            st.markdown(f"**🏷️ (直接放第二層)** `({len(no_c3_files)})`")
                        with col_d2_btn:
                            if st.button("📥", key=f"qup_{c1}_{c2}_none"):
                                st.session_state.quick_upload_target = (c1, c2, "")
                                st.rerun()
                        for f in no_c3_files:
                            if st.button(f"📄 {f}", key=f"dir_{c1}_{c2}_none_{f}", use_container_width=True):
                                st.session_state.jump_to_file = f

                        for c3 in list(c3_list):
                            c3_files = sorted([f for f, (fc1, fc2, fc3, _) in file_info_map.items() if fc1 == c1 and fc2 == c2 and fc3 == c3])
                            if not show_empty_dirs and not c3_files:
                                continue
                            col_d3_title, col_d3_edit, col_d3_up = st.columns([3.2, 0.9, 0.9])
                            with col_d3_title:
                                st.markdown(f"**🏷️ {c3}** `({len(c3_files)})`")
                            with col_d3_edit:
                                if st.button("✏️", key=f"edc3_{c1}_{c2}_{c3}"):
                                    st.session_state.edit_c3_target = (c1, c2, c3)
                                    st.rerun()
                            with col_d3_up:
                                if st.button("📥", key=f"qup_{c1}_{c2}_{c3}"):
                                    st.session_state.quick_upload_target = (c1, c2, c3)
                                    st.rerun()
                            for f in c3_files:
                                if st.button(f"📄 {f}", key=f"dir_{c1}_{c2}_{c3}_{f}", use_container_width=True):
                                    st.session_state.jump_to_file = f

    with st.expander("☁️ 專案雲端與 Figma 捷徑", expanded=True):
        cloud_links = db.get("cloud_links", DEFAULT_CLOUD_LINKS)
        for lk in cloud_links:
            st.markdown(f"[{lk['title']}]({lk['url']})")
        for fb in db.get("figma_boards", []):
            st.markdown(f"[🎨 {fb['title']} (開新視窗)]({fb['url']})")

        with st.popover("✏️ 管理捷徑連結", use_container_width=True):
            nl_title = st.text_input("連結顯示名稱：", key="nl_t")
            nl_url = st.text_input("網址 (URL)：", key="nl_u")
            if st.button("➕ 新增捷徑", type="primary", use_container_width=True) and nl_url.strip():
                clean_url = nl_url.strip()
                clean_title = nl_title.strip() or ("🎨 Figma 流程圖" if "figma.com" in clean_url.lower() else "🔗 外部雲端捷徑")
                db.setdefault("cloud_links", []).append({"title": clean_title, "url": clean_url})
                if "figma.com" in clean_url.lower():
                    db.setdefault("figma_boards", []).append({"title": clean_title, "url": clean_url})
                save_db(db)
                st.rerun()
            if cloud_links:
                st.divider()
                del_lk_idx = st.selectbox("選擇要移除的捷徑：", range(len(cloud_links)), format_func=lambda i: cloud_links[i]["title"], key="del_lk")
                if st.button("🗑️ 移除選中捷徑", use_container_width=True):
                    db["cloud_links"].pop(del_lk_idx)
                    save_db(db)
                    st.rerun()

    with st.expander("⚙️ 防呆上限與首頁快捷設定", expanded=False):
        new_max_mb = st.slider("單檔防呆上限 (MB)：", 20, 200, int(max_mb_limit), 10)
        if new_max_mb != max_mb_limit:
            db["max_mb"] = new_max_mb
            save_db(db)
            st.rerun()
        sc_str = st.text_area("自訂首頁快捷鈕（一行一個）：", value="\n".join(db.get("shortcuts", DEFAULT_SHORTCUTS)))
        if st.button("儲存快捷鈕", use_container_width=True):
            db["shortcuts"] = [s.strip() for s in sc_str.split("\n") if s.strip()]
            save_db(db)
            st.rerun()

    if raw_current_user == "聖引":
        with st.expander("👥 帳號管理 (聖引專屬)", expanded=False):
            new_u = st.text_input("➕ 新增工號/帳號：")
            new_name = st.text_input("設定中文姓名：")
            new_p = st.text_input("設定密碼：", value=TEAM_PASSWORD)
            if st.button("新增帳號", use_container_width=True) and new_u.strip():
                users_db[new_u.strip()] = {
                    "password": new_p.strip() or TEAM_PASSWORD, 
                    "role": "一般組員", 
                    "display_name": new_name.strip() or new_u.strip()
                }
                save_db(db)
                st.rerun()
            st.divider()
            ed_u = st.selectbox("選擇帳號：", list(users_db.keys()))
            ed_name = st.text_input("修改中文姓名：", value=users_db[ed_u].get("display_name", ed_u))
            ed_p = st.text_input("修改密碼：", value=users_db[ed_u]["password"])
            u1, u2 = st.columns(2)
            with u1:
                if st.button("更新設定", use_container_width=True):
                    users_db[ed_u]["password"] = ed_p.strip() or TEAM_PASSWORD
                    users_db[ed_u]["display_name"] = ed_name.strip() or ed_u
                    save_db(db)
                    st.rerun()
            with u2:
                if st.button("🗑️ 刪除帳號", use_container_width=True) and ed_u != "聖引":
                    del users_db[ed_u]
                    save_db(db)
                    st.rerun()

    with st.expander("💾 手動備份與還原", expanded=False):
        if os.path.exists(DB_FILE):
            with open(DB_FILE, "rb") as f:
                st.download_button("📥 下載備份 (.pkl)", data=f, file_name="engineering_knowledge_db.pkl", use_container_width=True)
        restore_file = st.file_uploader("📤 上傳舊有備份檔還原：", type=["pkl"])
        if restore_file and st.button("⚡ 立即還原並同步至 GitHub", use_container_width=True):
            with open(DB_FILE, "wb") as f:
                f.write(restore_file.getvalue())
            db = load_db()
            save_db(db)
            st.rerun()

    enable_web_search = st.checkbox("🌐 Google 聯網搜尋", value=False)

col_title, col_filter = st.columns([3, 2])
with col_title:
    st.subheader("🏗️ CV02 工程助手")
with col_filter:
    with st.popover("🎯 篩選檢索目錄範圍（預設全選）", use_container_width=True):
        selected_cat1s = st.multiselect("1. 第一層 (主夾)：", list(folder_tree.keys()), default=list(folder_tree.keys()))
        avail_c2 = [c2 for c1 in selected_cat1s for c2 in folder_tree.get(c1, {}).keys()]
        selected_cat2s = st.multiselect("2. 第二層 (次夾，留空代表全選)：", avail_c2, default=[])
        avail_c3 = [
            c3 for c1 in selected_cat1s
            for c2, c3_list in folder_tree.get(c1, {}).items()
            if (not selected_cat2s or c2 in selected_cat2s)
            for c3 in c3_list
        ]
        selected_cat3s = st.multiselect("3. 第三層 (細夾，留空代表全選)：", list(dict.fromkeys(avail_c3)), default=[])

if st.session_state.jump_to_file and st.session_state.jump_to_file in file_info_map:
    j_file = st.session_state.jump_to_file
    jc1, jc2, jc3, jup = file_info_map[j_file]
    with st.container(border=True):
        jh1, jh2 = st.columns([5, 1])
        with jh1:
            path_str = f"{jc1} > {jc2}" + (f" > {jc3}" if jc3 else "")
            st.markdown(f"### 📖 快速閱覽：`{j_file}`  `({path_str}｜上傳：{jup})`")
        with jh2:
            if st.button("✖️ 關閉閱覽", use_container_width=True):
                st.session_state.jump_to_file = None
                st.rerun()

        with st.expander("✏️ 就地改名 / 換資料夾 / 刪除此檔案", expanded=False):
            q_col1, q_col2, q_col3, q_col4 = st.columns([2, 2, 2, 2])
            with q_col1:
                q_new_name = st.text_input("檔案名稱：", value=j_file, key="q_ren")
            with q_col2:
                c1_opts = list(folder_tree.keys())
                q_c1 = st.selectbox("移至第1層：", c1_opts, index=c1_opts.index(jc1) if jc1 in c1_opts else 0, key="q_c1")
            with q_col3:
                c2_opts = list(folder_tree[q_c1].keys())
                q_c2 = st.selectbox("移至第2層：", c2_opts, index=c2_opts.index(jc2) if jc2 in c2_opts else 0, key="q_c2")
            with q_col4:
                raw_q3_opts = folder_tree[q_c1][q_c2]
                q3_opts_all = ["(無細夾 / 直接放第二層)"] + (raw_q3_opts if isinstance(raw_q3_opts, list) else [])
                curr_q3_idx = q3_opts_all.index(jc3) if jc3 in q3_opts_all else 0
                q_c3_sel = st.selectbox("移至第3層：", q3_opts_all, index=curr_q3_idx, key="q_c3_sel")
                q_c3 = "" if q_c3_sel == "(無細夾 / 直接放第二層)" else q_c3_sel

            qb1, qb2 = st.columns(2)
            with qb1:
                if st.button("💾 儲存變更", key="q_save", type="primary", use_container_width=True):
                    clean_qn = q_new_name.strip() or j_file
                    for c in chunks:
                        if c["file_name"] == j_file:
                            c["file_name"] = clean_qn
                            c["cat1"], c["cat2"], c["cat3"] = q_c1, q_c2, q_c3
                    save_db(db)
                    st.session_state.jump_to_file = clean_qn
                    st.rerun()
            with qb2:
                if st.button("🗑️ 刪除此檔案", key="q_del", use_container_width=True):
                    db["chunks"] = [c for c in chunks if c["file_name"] != j_file]
                    save_db(db)
                    st.session_state.jump_to_file = None
                    st.rerun()

        for fc in [c for c in chunks if c["file_name"] == j_file]:
            with st.expander(f"📄 {fc['location']}", expanded=True):
                st.text(fc["content"])

# ================= 核心六大分頁（新增「📊 施工進度甘特圖」） =================
tab_ai, tab_gantt, tab_figma, tab_memo, tab_search, tab_batch = st.tabs([
    "💬 AI 問答",
    "📊 施工進度甘特圖",
    "🎨 流程圖 (Figma)",
    "📝 公文生成",
    "🔎 搜檔案 & 讀全文",
    "📦 批次搬移與整理"
])

# ================= 💬 AI 問答 =================
with tab_ai:
    shortcuts = db.get("shortcuts", DEFAULT_SHORTCUTS)
    triggered_query = None
    if shortcuts:
        sc_cols = st.columns(len(shortcuts))
        for idx, sc_text in enumerate(shortcuts):
            if sc_cols[idx].button(sc_text, key=f"sc_{idx}", use_container_width=True):
                triggered_query = sc_text

    if "messages" not in st.session_state:
        st.session_state.messages = [
            {"role": "assistant", "content": f"**{current_user}** 您好！請直接在下方對話框打字、貼上文字、按 `Ctrl+V` 貼上螢幕截圖，或點擊 `➕` 夾帶檔案提問！"}
        ]

    for msg in st.session_state.messages:
        with st.chat_message(msg["role"]):
            st.markdown(msg["content"])

    if st.session_state.get("last_chat_files"):
        with st.container(border=True):
            col_arch1, col_arch2 = st.columns([3.5, 1.5])
            with col_arch1:
                f_names = "、".join([f.name for f in st.session_state.last_chat_files])
                st.markdown(f"💡 **剛剛隨附了臨時檔案：** `{f_names}`，需要直接存入永久目錄嗎？")
            with col_arch2:
                if st.button("💾 AI 自動分類存入永久目錄", type="primary", use_container_width=True):
                    if process_and_save_uploads(st.session_state.last_chat_files, mode="smart", enable_ocr=True):
                        st.session_state.last_chat_files = []
                        st.rerun()

    inline_files = []
    user_query = None
    try:
        chat_val = st.chat_input(
            "請輸入問題、直接貼上文字，或按 Ctrl+V 貼上截圖 / 點擊 ➕ 夾帶檔案...",
            accept_file="multiple",
            file_type=["pdf", "docx", "pptx", "xlsx", "txt", "jpg", "jpeg", "png"]
        )
        if chat_val:
            if isinstance(chat_val, str):
                user_query = chat_val
            else:
                user_query = getattr(chat_val, "text", "") or ""
                inline_files = list(getattr(chat_val, "files", []) or [])
                if not user_query and inline_files:
                    user_query = "請幫我分析隨附的圖片或檔案內容，並與內部工程資料庫比對。"
    except TypeError:
        chat_val = st.chat_input("請輸入問題（如：詳細價目表有沒有設計舊樁拔除、管子、里程碑分項...）")
        if chat_val:
            user_query = chat_val

    if triggered_query:
        user_query = triggered_query

    if user_query:
        st.session_state.last_chat_files = inline_files if inline_files else []
        display_q = f"{user_query} *(📎 隨附 {len(inline_files)} 份檔案/截圖)*" if inline_files else user_query

        st.session_state.messages.append({"role": "user", "content": display_q})
        with st.chat_message("user"):
            st.markdown(display_q)

        with st.chat_message("assistant"):
            with st.spinner("🔍 正在深度檢索工程資料庫中..."):
                internal_context = search_internal_db(user_query, db["chunks"], selected_cat1s, selected_cat2s, selected_cat3s)
                file_summary = [
                    f"[{c1}>{c2}{'>' + c3 if c3 else ''}] {fname}"
                    for fname, (c1, c2, c3, _) in file_info_map.items()
                    if c1 in selected_cat1s
                    and (not selected_cat2s or c2 in selected_cat2s)
                    and (not selected_cat3s or (not c3 or c3 in selected_cat3s))
                ]

                temp_text_context = ""
                multimodal_parts = []
                for tf in inline_files:
                    lname = tf.name.lower()
                    if lname.endswith((".jpg", ".jpeg", ".png")):
                        mime = "image/png" if lname.endswith(".png") else "image/jpeg"
                        multimodal_parts.append(types.Part.from_bytes(data=tf.getvalue(), mime_type=mime))
                        temp_text_context += f"\n【隨附影像/截圖：{tf.name}】\n"
                    elif lname.endswith(".pdf"):
                        t_chunks, _ = extract_text_from_file(tf, "臨時", "臨時", "", current_user, enable_ocr=False)
                        if t_chunks:
                            joined_t = "\n".join([f"{c['location']} {c['content']}" for c in t_chunks])
                            temp_text_context += f"\n=== 隨附臨時檔案：{tf.name} ===\n{joined_t[:15000]}\n"
                        else:
                            multimodal_parts.append(types.Part.from_bytes(data=tf.getvalue(), mime_type="application/pdf"))
                            temp_text_context += f"\n【隨附掃描PDF：{tf.name}】\n"
                    else:
                        t_chunks, _ = extract_text_from_file(tf, "臨時", "臨時", "", current_user, enable_ocr=False)
                        joined_t = "\n".join([f"{c['location']} {c['content']}" for c in t_chunks])
                        temp_text_context += f"\n=== 隨附臨時檔案：{tf.name} ===\n{joined_t[:15000]}\n"

                system_prompt = f"""
你是一位專業、嚴謹的工程專案 AI 顧問。請根據以下【內部工程資料庫檢索結果】與【臨時隨附資料】回答問題。

【目前檢索範圍內的文件清單】：
{", ".join(file_summary) if file_summary else "尚未匯入文件"}

【回答格式與排版鐵則（請嚴格遵守）】：
1. 🚫 嚴禁開頭廢話與重複問題：第一行請直接給出核心結論或答案！
2. 🚫 嚴禁結尾客套話。
3. 📚 參考依據與出處統一置底：請在整篇回答的最下方獨立列出「**📚 參考依據與出處：**」區塊。
4. 嚴禁瞎掰：若資料庫未載明，請直接告知「目前內部文件未載明」。

{temp_text_context}

=== 內部工程資料庫檢索結果 ===
{internal_context}
==============================

使用者問題：{user_query}
"""
                client = genai.Client(api_key=GEMINI_API_KEY)
                answer = None
                last_error = ""
                contents_payload = multimodal_parts + [system_prompt] if multimodal_parts else system_prompt

                for model_name in AVAILABLE_MODELS:
                    try:
                        config_args = {"temperature": 0.2}
                        if enable_web_search:
                            config_args["tools"] = [{"google_search": {}}]
                        response = client.models.generate_content(
                            model=model_name,
                            contents=contents_payload,
                            config=types.GenerateContentConfig(**config_args)
                        )
                        answer = response.text
                        break
                    except Exception as e:
                        last_error = str(e)
                        continue

                if not answer:
                    answer = f"⚠️ 呼叫失敗，請稍後再試。({last_error})"

                st.markdown(answer)
                st.session_state.messages.append({"role": "assistant", "content": answer})
                st.rerun()

# ================= 📊 施工進度甘特圖（開會互動專區：左表雙擊修改 ＋ 右圖動態縮放 ＋ Excel/PDF 匯出） =================
with tab_gantt:
    gantt_list = db.get("gantt_data", DEFAULT_GANTT_DATA)
    gantt_title = db.get("gantt_title", "台電統包工程 — 隧道明挖工法施工進度甘特圖與排程管控表")

    # 頂部開會儀表板與標題編輯
    g_top1, g_top2 = st.columns([3.5, 1.5])
    with g_top1:
        new_g_title = st.text_input("📌 專案排程表名稱（開會報告標題）：", value=gantt_title)
        if new_g_title.strip() != gantt_title:
            db["gantt_title"] = new_g_title.strip()
            save_db(db)
            st.rerun()
    with g_top2:
        with st.popover("📤 匯入 Excel 排程表 / 重置範本", use_container_width=True):
            st.markdown("##### 📥 從 Excel 匯入更新排程")
            st.caption("支援將本系統匯出的 `.xlsx` 修改後重新丟回，甘特圖將自動更新！")
            up_gantt_xlsx = st.file_uploader("選擇排程 Excel (.xlsx)：", type=["xlsx"], key="up_gantt_excel")
            if up_gantt_xlsx and st.button("⚡ 立即覆蓋匯入排程", type="primary", use_container_width=True):
                try:
                    df_in = pd.read_excel(up_gantt_xlsx, header=2)
                    new_g_rows = []
                    for idx, row in df_in.iterrows():
                        w_name = str(row.get("工項名稱", "")).strip()
                        if not w_name or w_name == "nan":
                            continue
                        s_str = str(row.get("預計開始日", "2026-11-01"))[:10]
                        e_str = str(row.get("預計完成日", "2026-11-10"))[:10]
                        try:
                            d_days = max(1, (pd.to_datetime(e_str) - pd.to_datetime(s_str)).days + 1)
                        except Exception:
                            d_days = int(row.get("工期(天)", 10))
                        cp_raw = str(row.get("是否要徑", ""))
                        is_cp_val = ("是" in cp_raw) or ("要徑" in cp_raw) or (cp_raw.lower() == "true")
                        new_g_rows.append({
                            "項次": str(row.get("項次", idx + 1)),
                            "WBS": str(row.get("WBS", f"{idx+1}.0")),
                            "工項名稱": w_name,
                            "預計開始日": s_str,
                            "預計完成日": e_str,
                            "工期(天)": d_days,
                            "進度(%)": int(float(row.get("進度(%)", 0) or 0)),
                            "狀態": str(row.get("狀態", "未開始")),
                            "是否要徑": is_cp_val,
                            "備註": str(row.get("備註與工法說明", row.get("備註", "")))
                        })
                    if new_g_rows:
                        db["gantt_data"] = new_g_rows
                        save_db(db)
                        st.success(f"✅ 已成功匯入 {len(new_g_rows)} 筆工項！")
                        st.rerun()
                except Exception as ex:
                    st.error(f"匯入失敗，請確認欄位格式：{ex}")
            st.divider()
            if st.button("🔄 還原為「隧道明挖工法」預設範本", use_container_width=True):
                db["gantt_data"] = [dict(x) for x in DEFAULT_GANTT_DATA]
                save_db(db)
                st.rerun()

    # 計算關鍵統計指標
    total_tasks = len(gantt_list)
    cp_tasks = sum(1 for r in gantt_list if r.get("是否要徑", False))
    done_tasks = sum(1 for r in gantt_list if r.get("狀態") == "已完成" or int(r.get("進度(%)", 0)) >= 100)
    try:
        min_start = min(pd.to_datetime([r["預計開始日"] for r in gantt_list])).strftime("%Y-%m-%d")
        max_end = max(pd.to_datetime([r["預計完成日"] for r in gantt_list])).strftime("%Y-%m-%d")
        total_span_days = (pd.to_datetime(max_end) - pd.to_datetime(min_start)).days + 1
    except Exception:
        min_start, max_end, total_span_days = "-", "-", 0

    m1, m2, m3, m4 = st.columns(4)
    m1.metric("📋 總工項數", f"{total_tasks} 項", f"已完成 {done_tasks} 項")
    m2.metric("🔥 關鍵要徑 (Critical Path)", f"{cp_tasks} 項", "重點控管工期")
    m3.metric("📅 專案預計起訖", f"{min_start} ～ {max_end}")
    m4.metric("⏳ 總日曆天數跨度", f"{total_span_days} 天")

    # 篩選與檢視控制列
    f_col1, f_col2, f_col3, f_col4 = st.columns([1.8, 1.8, 1.4, 1.5])
    with f_col1:
        filter_cp_only = st.checkbox("🔥 只顯示關鍵要徑 (Critical Path)", value=False)
    with f_col2:
        status_opts = ["全部狀態", "進行中", "未開始", "已完成", "進度落後"]
        filter_status = st.selectbox("🔍 狀態過濾：", status_opts, label_visibility="collapsed")
    with f_col3:
        color_mode = st.radio("🎨 圖表著色依據：", ["依要徑與狀態", "依完成進度(%)"], horizontal=True, label_visibility="collapsed")
    with f_col4:
        show_today_line = st.checkbox("📍 顯示今日基準紅線", value=True)

    # 過濾顯示用資料
    filtered_gantt = []
    for r in gantt_list:
        if filter_cp_only and not r.get("是否要徑", False):
            continue
        if filter_status != "全部狀態" and r.get("狀態") != filter_status:
            continue
        filtered_gantt.append(r)

    # 繪製 Plotly 互動式甘特圖
    fig_html_for_print = ""
    if filtered_gantt:
        df_plot = pd.DataFrame(filtered_gantt)
        df_plot["Start"] = pd.to_datetime(df_plot["預計開始日"], errors="coerce")
        df_plot["Finish"] = pd.to_datetime(df_plot["預計完成日"], errors="coerce") + pd.Timedelta(days=1)
        df_plot["Y軸標籤"] = df_plot.apply(
            lambda x: f"{'🔥 ' if x.get('是否要徑') else ''}[{x.get('WBS','')}] {x.get('工項名稱','')} ({x.get('工期(天)',0)}天)",
            axis=1
        )
        df_plot["類別顯示"] = df_plot.apply(
            lambda x: "🔥 關鍵要徑 (Critical Path)" if x.get("是否要徑") else f"🟢 一般工項 ({x.get('狀態','未開始')})",
            axis=1
        )
        df_plot["條上文字"] = df_plot.apply(
            lambda x: f"{x.get('工期(天)',0)}天 ({x.get('進度(%)',0)}%)",
            axis=1
        )

        if color_mode == "依要徑與狀態":
            color_map = {
                "🔥 關鍵要徑 (Critical Path)": "#D93025",
                "🟢 一般工項 (已完成)": "#2E7D32",
                "🟢 一般工項 (進行中)": "#1976D2",
                "🟢 一般工項 (未開始)": "#546E7A",
                "🟢 一般工項 (進度落後)": "#F57C00"
            }
            fig = px.timeline(
                df_plot,
                x_start="Start",
                x_end="Finish",
                y="Y軸標籤",
                color="類別顯示",
                color_discrete_map=color_map,
                text="條上文字",
                hover_name="工項名稱",
                hover_data={
                    "Start": False, "Finish": False, "Y軸標籤": False,
                    "WBS": True, "預計開始日": True, "預計完成日": True,
                    "工期(天)": True, "進度(%)": True, "狀態": True, "備註": True
                }
            )
        else:
            fig = px.timeline(
                df_plot,
                x_start="Start",
                x_end="Finish",
                y="Y軸標籤",
                color="進度(%)",
                color_continuous_scale="RdYlGn",
                range_color=[0, 100],
                text="條上文字",
                hover_name="工項名稱",
                hover_data={
                    "Start": False, "Finish": False, "Y軸標籤": False,
                    "WBS": True, "預計開始日": True, "預計完成日": True,
                    "工期(天)": True, "進度(%)": True, "狀態": True, "備註": True
                }
            )

        fig.update_yaxes(autorange="reversed", title="")
        fig.update_xaxes(
            title="工程時程軸 (可框選局部放大、雙擊恢復全貌)",
            tickformat="%Y/%m/%d",
            showgrid=True,
            gridcolor="#E0E0E0"
        )
        fig.update_traces(textposition="inside", insidetextanchor="middle")
        chart_h = max(420, len(df_plot) * 42 + 110)
        fig.update_layout(
            height=chart_h,
            margin=dict(l=10, r=20, t=35, b=35),
            legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="right", x=1, title=""),
            font=dict(family="Microsoft JhengHei, sans-serif", size=13),
            plot_bgcolor="#FFFFFF"
        )

        if show_today_line:
            today_str = date.today().strftime("%Y-%m-%d")
            fig.add_vline(x=today_str, line_width=2, line_dash="dash", line_color="#D93025")
            fig.add_annotation(
                x=today_str, y=1.02, yref="paper",
                text=f"📍 今日 ({today_str})", showarrow=False,
                font=dict(color="#D93025", size=12, weight="bold")
            )

        st.plotly_chart(fig, use_container_width=True, config={"displaylogo": False, "toImageButtonOptions": {"format": "png", "filename": "工程甘特圖", "scale": 2}})
        fig_html_for_print = fig.to_html(full_html=False, include_plotlyjs="cdn")
    else:
        st.warning("目前篩選條件下無對應工項。")

    # 一鍵匯出工具列（Excel / 開會列印版橫向 PDF 報表）
    dl_col1, dl_col2 = st.columns(2)
    with dl_col1:
        excel_bytes = generate_gantt_excel_bytes(gantt_list, gantt_title)
        st.download_button(
            label="📥 匯出精美 Excel 進度排程表 (.xlsx)",
            data=excel_bytes,
            file_name=f"{gantt_title}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            type="primary",
            use_container_width=True
        )
    with dl_col2:
        printable_html_bytes = generate_printable_gantt_html(gantt_list, gantt_title, fig_html_for_print)
        st.download_button(
            label="🖨️ 匯出開會用橫向 PDF 報表 (下載後點開即印)",
            data=printable_html_bytes,
            file_name=f"{gantt_title}_開會報表.html",
            mime="text/html",
            use_container_width=True
        )

    st.divider()
    st.markdown("#### ✏️ 開會現場即時排程編輯表（雙擊格子即可修改日期、進度、勾選要徑，亦可於最下方新增工項列）")
    st.caption("💡 修改「預計開始日」或「預計完成日」後，點擊下方 **「💾 儲存排程變更並重算甘特圖」**，系統會自動計算工期天數並更新上方圖表！")

    df_edit = pd.DataFrame(gantt_list)
    df_edit["預計開始日"] = pd.to_datetime(df_edit["預計開始日"], errors="coerce").dt.date
    df_edit["預計完成日"] = pd.to_datetime(df_edit["預計完成日"], errors="coerce").dt.date

    edited_gantt_df = st.data_editor(
        df_edit,
        num_rows="dynamic",
        column_config={
            "項次": st.column_config.TextColumn("項次", width="small"),
            "WBS": st.column_config.TextColumn("WBS", width="small"),
            "工項名稱": st.column_config.TextColumn("📌 工項名稱 (雙擊可改)", width="medium", required=True),
            "預計開始日": st.column_config.DateColumn("📅 預計開始日", format="YYYY-MM-DD", required=True),
            "預計完成日": st.column_config.DateColumn("📅 預計完成日", format="YYYY-MM-DD", required=True),
            "工期(天)": st.column_config.NumberColumn("⏳ 工期(天) [自動重算]", disabled=True, width="small"),
            "進度(%)": st.column_config.NumberColumn("📈 進度(%)", min_value=0, max_value=100, step=5, format="%d%%", width="small"),
            "狀態": st.column_config.SelectboxColumn("狀態", options=["未開始", "進行中", "已完成", "進度落後"], width="small"),
            "是否要徑": st.column_config.CheckboxColumn("🔥 關鍵要徑", default=False, width="small"),
            "備註": st.column_config.TextColumn("📝 備註與工法說明", width="large")
        },
        hide_index=True,
        use_container_width=True,
        key="gantt_table_editor"
    )

    if st.button("💾 儲存排程變更並重算甘特圖 (同步備份至 GitHub)", type="primary", use_container_width=True):
        updated_list = []
        for idx, row in edited_gantt_df.iterrows():
            w_name = str(row.get("工項名稱", "")).strip()
            if not w_name or w_name == "None":
                continue
            s_d = row.get("預計開始日")
            e_d = row.get("預計完成日")
            s_str = s_d.strftime("%Y-%m-%d") if pd.notnull(s_d) and hasattr(s_d, "strftime") else "2026-11-01"
            e_str = e_d.strftime("%Y-%m-%d") if pd.notnull(e_d) and hasattr(e_d, "strftime") else s_str
            try:
                calc_days = max(1, (pd.to_datetime(e_str) - pd.to_datetime(s_str)).days + 1)
            except Exception:
                calc_days = 1
            prog_v = int(row.get("進度(%)", 0)) if pd.notnull(row.get("進度(%)")) else 0
            stat_v = str(row.get("狀態", "未開始")) if pd.notnull(row.get("狀態")) else "未開始"
            if prog_v >= 100:
                stat_v = "已完成"
            updated_list.append({
                "項次": str(row.get("項次") or (len(updated_list) + 1)),
                "WBS": str(row.get("WBS") or f"{len(updated_list) + 1}.0"),
                "工項名稱": w_name,
                "預計開始日": s_str,
                "預計完成日": e_str,
                "工期(天)": calc_days,
                "進度(%)": prog_v,
                "狀態": stat_v,
                "是否要徑": bool(row.get("是否要徑", False)),
                "備註": str(row.get("備註") or "")
            })
        db["gantt_data"] = updated_list
        with st.spinner("💾 正在重算工期並同步至 GitHub..."):
            save_db(db)
        st.success("✅ 排程與甘特圖已更新！")
        st.rerun()

# ================= 🎨 流程圖 (Figma) =================
with tab_figma:
    f_head1, f_head2 = st.columns([4, 1.5])
    with f_head1:
        st.markdown("### 🎨 專案工程流程圖與架構看板 (Figma 互動內嵌)")
    with f_head2:
        with st.popover("➕ 新增 / 管理 Figma 流程圖", use_container_width=True):
            fg_title = st.text_input("流程圖名稱：", placeholder="如：施工查驗與送審流程圖")
            fg_url = st.text_input("Figma 分享連結或 Embed 語法：", placeholder="https://www.figma.com/...")
            if st.button("💾 儲存並內嵌此流程圖", type="primary", use_container_width=True):
                if fg_url.strip():
                    clean_fg_title = fg_title.strip() or "🎨 專案工程流程圖"
                    db.setdefault("figma_boards", []).append({"title": clean_fg_title, "url": fg_url.strip()})
                    save_db(db)
                    st.rerun()
            figma_list = db.get("figma_boards", [])
            if figma_list:
                st.divider()
                del_fg_idx = st.selectbox("選擇要移除的項目：", range(len(figma_list)), format_func=lambda i: figma_list[i]["title"])
                if st.button("🗑️ 刪除此流程圖", use_container_width=True):
                    db["figma_boards"].pop(del_fg_idx)
                    save_db(db)
                    st.rerun()

    figma_boards = db.get("figma_boards", [])
    if not figma_boards:
        st.info("👆 目前尚未加入 Figma 流程圖！請點擊右上角的 **「➕ 新增 / 管理 Figma 流程圖」** 貼上網址即可！")
    else:
        ctrl1, ctrl2, ctrl3 = st.columns([2.2, 1.3, 2.0])
        with ctrl1:
            sel_fg_idx = st.selectbox("📂 選擇看板：", range(len(figma_boards)), format_func=lambda i: f"🎨 {figma_boards[i]['title']}")
        selected_board = figma_boards[sel_fg_idx]
        with ctrl2:
            canvas_height = st.slider("📐 畫布高度 (px)：", 500, 1200, 750, 50)
        with ctrl3:
            st.write("")
            raw_link = selected_board["url"]
            if "iframe" in raw_link and "src=" in raw_link:
                m = re.search(r'src=["\']([^"\']+)["\']', raw_link)
                raw_link = m.group(1) if m else raw_link
            st.link_button("🚀 畫面讀不到？點此直接開 Figma 新視窗", raw_link, type="primary", use_container_width=True)

        embed_src = to_figma_embed_url(selected_board["url"])
        components.iframe(embed_src, height=canvas_height, scrolling=True)

# ================= 📝 公文生成（已修復多模型自動備援，不會再跳 ServerError） =================
with tab_memo:
    st.markdown("### 📝 智慧公文生成與台電標準簽辦用紙排版室")
    st.caption("支援一鍵匯出 **「台灣電力公司 簽辦用紙（紅框內簽標準表格）」** 或標準外發函 Word (.docx) 檔！")

    memo_topic = st.text_area("📌 請輸入擬稿需求或事件摘要：", height=100)
    m_col1, m_col2 = st.columns(2)
    with m_col1:
        memo_type = st.selectbox("📄 文別與版型選擇：", [
            "台電標準內簽 (台灣電力公司 簽辦用紙 紅框版)",
            "外發函 (標準公文函)",
            "書函",
            "開會通知單",
            "會議紀錄"
        ])
    with m_col2:
        memo_target = st.text_input("🏢 受文對象 / 陳核層級：", placeholder="例如：吉興工程顧問有限公司 / 或 組長、處長")

    unit_name_val, doc_no_val, file_no_val, retention_val = "水工組", "", "", ""
    if "簽辦用紙" in memo_type:
        u_c1, u_c2, u_c3, u_c4 = st.columns(4)
        with u_c1:
            unit_name_val = st.text_input("🏷️ 簽辦單位：", value="水工組")
        with u_c2:
            doc_no_val = st.text_input("🔢 文號 (選填)：", value="")
        with u_c3:
            file_no_val = st.text_input("📁 檔號 (選填)：", value="")
        with u_c4:
            retention_val = st.text_input("⏳ 保存年限 (選填)：", value="10年")

    if st.button("🚀 開始智慧擬稿與排版", type="primary", use_container_width=True):
        if not memo_topic.strip():
            st.warning("請先輸入擬稿需求！")
        else:
            with st.spinner("🤖 AI 正在深度參考您的歷史公文風格、契約規範進行擬稿中..."):
                memo_history_chunks = [
                    c for c in db["chunks"] 
                    if "公文" in c.get("cat1", "") or "公文" in c.get("cat2", "") or "往來" in c.get("cat2", "")
                ]
                history_text = "\n".join([f"【參考範本：{c['file_name']}】\n{c['content']}" for c in memo_history_chunks[:10]])
                contract_context = search_internal_db(memo_topic, db["chunks"], selected_cat1s, selected_cat2s, selected_cat3s)

                is_internal_sign = "簽辦用紙" in memo_type
                format_instruction = (
                    "本次為【台灣電力公司 簽辦用紙（內部簽呈）】，請依序輸出：主旨、說明（一、二、三...）、擬辦、陳核。"
                    if is_internal_sign else
                    "請嚴格遵守標準工程公文格式，包含【受文者】、【主旨】、【說明】（分項條列一、二、三）、【辦法】。"
                )

                memo_prompt = f"""
你是一位專業、嚴謹的台灣電力公司資深工程師。擬稿人：{current_user}（簽辦單位：{unit_name_val}）。
請根據以下資訊產出一份標準工程公文。

【文別】：{memo_type}
【受文/陳核對象】：{memo_target if memo_target.strip() else "依台電工程公文慣例"}
【擬稿核心需求】：{memo_topic}

【歷史公文參考範本】：
{history_text if history_text.strip() else "（請依標準台電工程公文格式產出）"}

【內部合約與資料庫檢索結果】：
{contract_context}

【排版與內容鐵則】：
1. 🚫 嚴禁任何 AI 開頭問候語或結尾客套話，第一行直接從公文正題開始輸出！
2. {format_instruction}
"""
                client = genai.Client(api_key=GEMINI_API_KEY)
                generated_memo_text = None
                last_err = ""
                for m_name in AVAILABLE_MODELS:
                    try:
                        memo_res = client.models.generate_content(
                            model=m_name,
                            contents=memo_prompt,
                            config=types.GenerateContentConfig(temperature=0.2)
                        )
                        if memo_res and memo_res.text:
                            generated_memo_text = memo_res.text
                            break
                    except Exception as e:
                        last_err = str(e)
                        continue

                if not generated_memo_text:
                    generated_memo_text = f"⚠️ 公文生成伺服器暫時忙碌，請再按一次生成按鈕試試！({last_err})"

                st.session_state.last_generated_memo = generated_memo_text
                st.session_state.last_memo_meta = {
                    "is_internal": is_internal_sign,
                    "unit": unit_name_val,
                    "doc_no": doc_no_val,
                    "file_no": file_no_val,
                    "retention": retention_val
                }
                st.rerun()

    if st.session_state.get("last_generated_memo"):
        st.divider()
        edited_memo = st.text_area("✏️ 您可在此直接微調內文，下方按鈕會即時打包成標準 Word 檔：", value=st.session_state.last_generated_memo, height=320)
        st.session_state.last_generated_memo = edited_memo

        meta = st.session_state.get("last_memo_meta", {"is_internal": True, "unit": "水工組", "doc_no": "", "file_no": "", "retention": ""})
        if meta.get("is_internal", True):
            docx_bytes = generate_taipower_internal_memo_docx(
                unit_name=meta.get("unit", "水工組"),
                doc_no=meta.get("doc_no", ""),
                file_no=meta.get("file_no", ""),
                retention=meta.get("retention", ""),
                memo_text=edited_memo
            )
            dl_label = "📥 一鍵下載「台灣電力公司 簽辦用紙（紅框標準版）」Word 檔 (.docx)"
            dl_fname = f"台電簽辦用紙_{meta.get('unit', '水工組')}_{current_user}.docx"
        else:
            doc_gen = docx.Document()
            doc_gen.add_heading(f"工程公文草稿 — {current_user} 擬", level=1)
            for line in edited_memo.split("\n"):
                if line.strip().startswith("###") or line.strip().startswith("【"):
                    doc_gen.add_heading(line.replace("#", "").strip(), level=2)
                else:
                    doc_gen.add_paragraph(line)
            docx_buf = io.BytesIO()
            doc_gen.save(docx_buf)
            docx_bytes = docx_buf.getvalue()
            dl_label = "📥 一鍵下載排版完整的公文 Word 檔 (.docx)"
            dl_fname = f"工程公文草稿_{current_user}.docx"

        st.download_button(
            label=dl_label,
            data=docx_bytes,
            file_name=dl_fname,
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            type="primary",
            use_container_width=True
        )

# ================= 🔎 搜檔案 & 讀全文 =================
with tab_search:
    col_s1, col_s2 = st.columns([3, 1])
    with col_s1:
        search_kw = st.text_input("🔎 輸入關鍵字秒查檔案內文（多個詞用空白隔開）：")
    with col_s2:
        match_mode = st.radio("匹配模式：", ["包含任一 (OR)", "包含全部 (AND)"], horizontal=True)

    if search_kw.strip():
        raw_kws = [k.strip() for k in search_kw.strip().split() if k.strip()]
        matched_results = []
        for c in chunks:
            if c.get("cat1") not in selected_cat1s:
                continue
            if selected_cat2s and c.get("cat2") not in selected_cat2s:
                continue
            c3_val = c.get("cat3", "")
            if selected_cat3s and c3_val and c3_val not in selected_cat3s:
                continue
            target_text = f"{c['file_name']} {c['content']}".lower()
            if match_mode == "包含全部 (AND)":
                is_match = all(k.lower() in target_text for k in raw_kws)
            else:
                is_match = any(k.lower() in target_text for k in raw_kws)
            if is_match:
                hits = sum(target_text.count(k.lower()) for k in raw_kws)
                matched_results.append((hits, c))

        matched_results.sort(key=lambda x: x[0], reverse=True)
        st.success(f"🎯 共找到符合條件的段落總計 **{len(matched_results)}** 處！")

        file_grouped_hits = {}
        for hits, item in matched_results:
            fname = item["file_name"]
            if fname not in file_grouped_hits:
                file_grouped_hits[fname] = {
                    "cat1": item["cat1"], "cat2": item["cat2"], "cat3": item.get("cat3", ""),
                    "uploader": item.get("uploader", "系統"), "total_hits": 0, "chunks": []
                }
            file_grouped_hits[fname]["total_hits"] += hits
            file_grouped_hits[fname]["chunks"].append((hits, item))

        for fname, gdata in file_grouped_hits.items():
            c3_str = f" > {gdata['cat3']}" if gdata['cat3'] else ""
            card_title = f"📄 {fname} ｜ 🗂️ {gdata['cat1']} > {gdata['cat2']}{c3_str} ｜ 命中：{len(gdata['chunks'])} 處"
            with st.expander(card_title, expanded=False):
                view_mode = st.radio("檢視模式：", ["🎯 只看關鍵字命中段落", f"📖 閱讀《{fname}》完整全文"], horizontal=True, key=f"vmode_{fname}")
                if view_mode == "🎯 只看關鍵字命中段落":
                    for idx, (h_cnt, chk) in enumerate(gdata["chunks"], 1):
                        st.markdown(f"**📍 命中位置：`{chk['location']}`**")
                        hl_html = highlight_keywords(chk["content"], raw_kws).replace("\n", "<br>")
                        st.markdown(f"<div style='background-color: #f8f9fa; padding: 10px; border-radius: 6px; border-left: 3px solid #ffc107;'>{hl_html}</div>", unsafe_allow_html=True)
                else:
                    for fc in [c for c in chunks if c["file_name"] == fname]:
                        st.markdown(f"**📄 {fc['location']}**")
                        hl_full = highlight_keywords(fc["content"], raw_kws).replace("\n", "<br>")
                        st.markdown(f"<div style='background-color: #ffffff; padding: 10px; border-radius: 6px; border: 1px solid #e0e0e0; margin-bottom: 8px;'>{hl_full}</div>", unsafe_allow_html=True)

    st.divider()
    with st.expander("📖 直接從清單挑選檔案閱讀全文", expanded=not bool(search_kw.strip())):
        if not file_info_map:
            st.info("目前資料庫尚無檔案。")
        else:
            sorted_files = sorted(list(file_info_map.keys()))
            default_idx = sorted_files.index(st.session_state.jump_to_file) if st.session_state.jump_to_file in sorted_files else 0
            view_file = st.selectbox("選擇要閱讀完整全文的檔案：", sorted_files, index=default_idx)
            for fc in [c for c in chunks if c["file_name"] == view_file]:
                with st.expander(f"📄 {fc['file_name']} — {fc['location']}", expanded=True):
                    st.text(fc["content"])

# ================= 📦 批次搬移與整理 =================
with tab_batch:
    st.markdown("#### 📦 檔案批次搬移、改名與刪除總管")
    if not file_info_map:
        st.info("目前資料庫尚無檔案。")
    else:
        b_col1, b_col2 = st.columns([2, 2])
        with b_col1:
            filter_b_c1 = st.selectbox("🔍 篩選來源主夾：", ["(顯示全部檔案)"] + list(folder_tree.keys()))
        with b_col2:
            kw_filter_file = st.text_input("🔍 檔名關鍵字過濾：")

        display_files = sorted(list(file_info_map.keys()))
        if filter_b_c1 != "(顯示全部檔案)":
            display_files = [f for f in display_files if file_info_map[f][0] == filter_b_c1]
        if kw_filter_file.strip():
            display_files = [f for f in display_files if kw_filter_file.strip().lower() in f.lower()]

        table_data = []
        for f in display_files:
            fc1, fc2, fc3, fup = file_info_map[f]
            path_display = f"{fc1} > {fc2}" + (f" > {fc3}" if fc3 else " (直接放第二層)")
            table_data.append({
                "勾選": False,
                "檔案名稱 (雙擊可直接改名)": f,
                "目前所在三層路徑": path_display,
                "上傳者": fup,
                "_orig_name": f
            })

        edited_table = st.data_editor(
            table_data,
            column_config={
                "勾選": st.column_config.CheckboxColumn("☑️ 勾選", default=False),
                "檔案名稱 (雙擊可直接改名)": st.column_config.TextColumn("📄 檔案名稱 (雙擊可直接改名)", width="large"),
                "目前所在三層路徑": st.column_config.TextColumn("🗂️ 目前所在三層路徑", disabled=True, width="large"),
                "上傳者": st.column_config.TextColumn("👤 上傳者", disabled=True, width="small"),
                "_orig_name": None
            },
            hide_index=True,
            use_container_width=True,
            key="batch_editor"
        )

        selected_rows = [r for r in edited_table if r["勾選"]]
        renamed_rows = [r for r in edited_table if r["檔案名稱 (雙擊可直接改名)"].strip() != r["_orig_name"]]

        if renamed_rows:
            if st.button(f"✏️ 立即儲存 {len(renamed_rows)} 份已修改的檔案名稱", type="primary"):
                for r in renamed_rows:
                    old_n = r["_orig_name"]
                    new_n = r["檔案名稱 (雙擊可直接改名)"].strip()
                    if new_n:
                        for c in chunks:
                            if c["file_name"] == old_n:
                                c["file_name"] = new_n
                save_db(db)
                st.rerun()

        st.divider()
        m_col1, m_col2, m_col3, m_col4, m_col5 = st.columns([2, 2, 2, 2, 2])
        with m_col1:
            dest_c1 = st.selectbox("目標第1層：", list(folder_tree.keys()), key="b_c1")
        with m_col2:
            dest_c2 = st.selectbox("目標第2層：", list(folder_tree[dest_c1].keys()), key="b_c2")
        with m_col3:
            raw_b_c3 = folder_tree[dest_c1][dest_c2]
            b_c3_opts = ["(無細夾 / 直接放第二層)"] + (raw_b_c3 if isinstance(raw_b_c3, list) else [])
            b_c3_sel = st.selectbox("目標第3層：", b_c3_opts, key="b_c3_sel")
            dest_c3 = "" if b_c3_sel == "(無細夾 / 直接放第二層)" else b_c3_sel
        with m_col4:
            st.write("")
            if st.button(f"📦 一鍵搬移 ({len(selected_rows)} 份)", type="primary", use_container_width=True, disabled=(len(selected_rows) == 0)):
                sel_names = {r["_orig_name"] for r in selected_rows}
                for c in chunks:
                    if c["file_name"] in sel_names:
                        c["cat1"], c["cat2"], c["cat3"] = dest_c1, dest_c2, dest_c3
                save_db(db)
                st.rerun()
        with m_col5:
            st.write("")
            if st.button(f"🗑️ 批次刪除 ({len(selected_rows)} 份)", use_container_width=True, disabled=(len(selected_rows) == 0)):
                sel_names = {r["_orig_name"] for r in selected_rows}
                db["chunks"] = [c for c in chunks if c["file_name"] not in sel_names]
                save_db(db)
                st.rerun()
